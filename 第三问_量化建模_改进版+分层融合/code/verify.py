"""独立复核：不调用模型函数，直接用 CSV 账本重算结算、平衡与工作表一致性。"""
from pathlib import Path
import json
import numpy as np
import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
R = ROOT / 'results'


def main():
    iv = pd.read_csv(R / 'intervals.csv')
    daily = pd.read_csv(R / 'daily_6+12+18.csv')
    p = json.loads((R / 'parameters.json').read_text())['selected']
    pars = json.loads((R / 'parameters.json').read_text())

    assert len(iv) == 334 * 144 and len(daily) == 334
    assert iv.groupby('date').size().eq(144).all()
    q = iv.adjusted.to_numpy(); q0 = iv.plan.to_numpy()
    c = iv.charge.to_numpy(); d = iv.discharge.to_numpy(); z = iv.emergency.to_numpy(); s = iv.spill.to_numpy()
    e0 = iv.soc_start.to_numpy(); e1 = iv.soc_end.to_numpy(); price = iv.price.to_numpy()

    balance = q + z + d - c - s - (iv.load_kw - iv.pv_kw).to_numpy() / 6
    dynamics = e1 - e0 - p['eta'] * c + d / p['eta']

    # 改进后主口径的独立重算：取消部分只承担 50% 违约成本（减购为负费用）
    assert p['settlements'] == 'net_refund' and p['terminal'] == 'inventory_value'
    settlement = (price * q0 + price * p['up'] * np.maximum(q - q0, 0)
                  - price * p['down'] * np.maximum(q0 - q, 0) + price * p['emergency'] * z)

    assert np.max(abs(balance)) < 1e-7 and np.max(abs(dynamics)) < 1e-7
    assert np.min(e0) >= 1200 - 1e-7 and np.max(e1) <= 10800 + 1e-7
    assert max(c.max(), d.max()) <= 5000 / 6 + 1e-7 and np.minimum(c, d).max() < 1e-7
    assert np.max(abs(e1[:-1] - e0[1:])) < 1e-7
    assert np.min(q) >= -1e-7
    daycost = settlement.reshape(334, 144).sum(axis=1)
    assert np.max(abs(daycost - daily.total_cost)) < 1e-5

    # 分解项独立重算
    plan = (price * q0).reshape(334, 144).sum(axis=1)
    up = (price * p['up'] * np.maximum(q - q0, 0)).reshape(334, 144).sum(axis=1)
    dn = (-price * p['down'] * np.maximum(q0 - q, 0)).reshape(334, 144).sum(axis=1)
    em = (price * p['emergency'] * z).reshape(334, 144).sum(axis=1)
    for arr, col in [(plan, 'plan_cost'), (up, 'up_cost'), (dn, 'down_net_cost'), (em, 'emergency_cost')]:
        assert np.max(abs(arr - daily[col])) < 1e-5, col
    assert np.max(abs((plan + up + dn + em) - daily.total_cost)) < 1e-4

    # 库存校正：v_E = p_min/eta
    vE = float(price.min() / p['eta'])
    inv = daily.total_cost.to_numpy() + vE * (daily.soc_start.to_numpy() - daily.soc_end.to_numpy())
    assert np.max(abs(inv - daily.inventory_adjusted_cost)) < 1e-5

    # ---------------- 提交表一致性 ----------------
    w = openpyxl.load_workbook(ROOT / 'result3.xlsx', data_only=True)
    assert w.sheetnames == ['计划购电量', '调整购电量', '充放电量', '紧急购电量'], w.sheetnames
    for name, key, costkey in [('计划购电量', 'plan', 'plan_cost'), ('调整购电量', 'adjusted', 'total_cost')]:
        vals = list(w[name].values)
        assert len(vals) == 335 and len(vals[0]) == 147
        assert vals[0][1] == '00:00-00:10' and vals[0][144] == '23:50-24:00'
        assert vals[0][145] == '全天购电量' and vals[0][146] == '全天购电费'
        arr = np.array([x[1:145] for x in vals[1:]], float)
        expected = iv[key].to_numpy().reshape(334, 144)
        assert np.max(abs(arr - expected)) < 1e-7
        quantities = np.array([x[145] for x in vals[1:]], float)
        assert np.max(abs(quantities - expected.sum(axis=1))) < 1e-6
        costs = np.array([x[146] for x in vals[1:]], float)
        assert np.max(abs(costs - daily[costkey])) < 1e-6
        for row, ds in zip(vals[1:], daily.date):
            assert pd.Timestamp(row[0]).date() == pd.Timestamp(ds).date()
    charge = list(w['充放电量'].values)[1:]
    assert len(charge) == 334 * 6
    actual = np.array([x[2:4] for x in charge], float)
    expected = np.column_stack([c.reshape(334 * 6, 24).sum(axis=1), d.reshape(334 * 6, 24).sum(axis=1)])
    assert np.max(abs(actual - expected)) < 1e-6
    for i in range(334):
        assert abs(charge[6 * i][5] - daily.soc_start.iloc[i]) < 1e-6
        assert abs(charge[6 * i + 5][5] - daily.soc_end.iloc[i]) < 1e-6
        assert charge[6 * i + 5][4] == '24:00' and charge[6 * i][4] == '00:00'
    ev = list(w['紧急购电量'].values)[1:]
    assert abs(sum(row[2] for row in ev) - z.sum()) < 1e-6
    w.close()

    check = {'intervals': len(iv), 'days': len(daily),
             'independent_balance_max_error': float(abs(balance).max()),
             'independent_soc_max_error': float(abs(dynamics).max()),
             'independent_cost_max_daily_error': float(abs(daycost - daily.total_cost).max()),
             'independent_component_max_error': float(max(abs(plan - daily.plan_cost).max(), abs(up - daily.up_cost).max(),
                                                           abs(dn - daily.down_net_cost).max(), abs(em - daily.emergency_cost).max())),
             'vE_used': vE, 'declared_vE': pars['terminal_inventory_value_vE'],
             'xlsx_values_match_ledger': True, 'xlsx_summary_columns_match': True, 'xlsx_all_dates_match': True,
             'xlsx_charge_and_emergency_reconcile': True, 'xlsx_block_markers_fixed': True, 'passed': True}
    assert abs(check['vE_used'] - check['declared_vE']) < 1e-9
    (R / 'delivery_verification.json').write_text(json.dumps(check, indent=2, ensure_ascii=False))
    print(json.dumps(check, indent=2, ensure_ascii=False))


if __name__ == '__main__':
    main()
