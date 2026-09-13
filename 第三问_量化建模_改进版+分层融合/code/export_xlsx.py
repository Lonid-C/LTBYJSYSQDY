"""生成提交用 result3.xlsx（四个业务工作表）。

结构与原模板一致：日期 + 144 个 10 分钟交付时段 + 全天购电量 + 全天购电费（前两张表）。
时间表头明确校正为当天 00:00—24:00。最后两列为数值列（而非依赖第三方表格工具重算的公式），
以保证在任意查看器中都能直接显示，其数值由 code/verify.py 独立重算校验。
"""
import json
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
R = ROOT / 'results'
SELECTED = ['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']
HEAD_FILL = PatternFill('solid', fgColor='325B88')
HEAD_FONT = Font(name='Arial', size=10, bold=True, color='FFFFFF')
BODY_FONT = Font(name='Arial', size=10)
THIN = Side(style='thin', color='D8DEE6')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def clock(t):
    return f'{t // 6:02d}:{(t % 6) * 10:02d}'


def interval(t):
    return f'{clock(t)}-{clock(t + 1)}'


def style_header(ws, ncol):
    for j in range(1, ncol + 1):
        c = ws.cell(row=1, column=j)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'B2'
    ws.sheet_view.showGridLines = False


def autosize(ws, ncol, nrow, width=17.0, first_width=None):
    for j in range(1, ncol + 1):
        ws.column_dimensions[get_column_letter(j)].width = (first_width if j == 1 and first_width else width)


def main():
    iv = pd.read_csv(R / 'intervals.csv')
    daily = pd.read_csv(R / 'daily_6+12+18.csv')
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---------- 计划购电量 / 调整购电量 ----------
    header = ['日期\\时间'] + [interval(t) for t in range(144)] + ['全天购电量', '全天购电费']
    for sheet, key, costkey, note in [('计划购电量', 'plan', 'plan_cost', '0:00 制定的基准计划购电量'),
                                      ('调整购电量', 'adjusted', 'total_cost', '相对 0:00 基准计划净额结算后的最终有效购电量')]:
        ws = wb.create_sheet(sheet)
        ws.append(header)
        for date, g in iv.groupby('date', sort=True):
            g = g.sort_values('t')
            d = daily[daily.date == date].iloc[0]
            ws.append([date] + [float(x) for x in g[key].tolist()] + [float(g[key].sum()), float(d[costkey])])
        style_header(ws, len(header))
        autosize(ws, len(header), ws.max_row)
        for i in range(2, ws.max_row + 1):
            ws.cell(row=i, column=1).number_format = 'yyyy-mm-dd'
            ws.cell(row=i, column=1).alignment = Alignment(horizontal='center')
            ws.cell(row=i, column=1).font = BODY_FONT
            for j in range(2, len(header) + 1):
                cell = ws.cell(row=i, column=j)
                cell.number_format = '#,##0.000'
                cell.font = BODY_FONT
        ws.column_dimensions[get_column_letter(len(header) - 1)].width = 22
        ws.column_dimensions[get_column_letter(len(header))].width = 22

    # ---------- 充放电量 ----------
    ws = wb.create_sheet('充放电量')
    ws.append(['日期', '时间段', '充电量', '放电量', '时刻', '储电量'])
    for date, g in iv.groupby('date', sort=True):
        g = g.sort_values('t').reset_index(drop=True)
        for start in range(0, 144, 24):
            piece = g.iloc[start:start + 24]
            block = f'{clock(start)}-{clock(start + 24)}'
            mark = '00:00' if start == 0 else ('24:00' if start == 120 else '')
            soc = g.soc_start.iloc[0] if start == 0 else (g.soc_end.iloc[-1] if start == 120 else None)
            ws.append([date, block, float(piece.charge.sum()), float(piece.discharge.sum()), mark,
                       None if soc is None else float(soc)])
    style_header(ws, 6)
    for i in range(2, ws.max_row + 1):
        ws.cell(row=i, column=1).number_format = 'yyyy-mm-dd'
        ws.cell(row=i, column=6).number_format = '#,##0.000'
        for j in range(1, 7):
            ws.cell(row=i, column=j).font = BODY_FONT
        ws.cell(row=i, column=2).alignment = Alignment(horizontal='center')
        ws.cell(row=i, column=5).alignment = Alignment(horizontal='center')
    autosize(ws, 6, ws.max_row, first_width=12.0)
    ws.column_dimensions['B'].width = 22

    # ---------- 紧急购电量 ----------
    ws = wb.create_sheet('紧急购电量')
    ws.append(['日期', '购电时间段', '购电量'])
    for date, g in iv.groupby('date', sort=True):
        g = g.sort_values('t').reset_index(drop=True)
        j, found = 0, False
        while j < 144:
            if g.emergency.iloc[j] <= 1e-7:
                j += 1
                continue
            start = j
            while j < 144 and g.emergency.iloc[j] > 1e-7:
                j += 1
            ws.append([date, f'{clock(start)}-{clock(j)}', float(g.emergency.iloc[start:j].sum())])
            found = True
        if not found:
            ws.append([date, '无', 0.0])
    style_header(ws, 3)
    for i in range(2, ws.max_row + 1):
        ws.cell(row=i, column=1).number_format = 'yyyy-mm-dd'
        ws.cell(row=i, column=3).number_format = '#,##0.000'
        for j in range(1, 4):
            ws.cell(row=i, column=j).font = BODY_FONT
        ws.cell(row=i, column=2).alignment = Alignment(horizontal='center')
    autosize(ws, 3, ws.max_row, first_width=12.0)
    ws.column_dimensions['B'].width = 24

    out = ROOT / 'result3.xlsx'
    wb.save(out)
    (R / 'export.log').write_text(
        f'sheets={wb.sheetnames}\nrows={[wb[s].max_row for s in wb.sheetnames]}\ncols={[wb[s].max_column for s in wb.sheetnames]}\n'
        f'selected_dates_present={sorted(set(SELECTED) & set(iv.date.unique()))}\n', encoding='utf8')
    print('saved', out, wb.sheetnames)


if __name__ == '__main__':
    main()
