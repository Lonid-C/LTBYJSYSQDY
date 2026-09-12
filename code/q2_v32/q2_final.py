"""Q2 V3.2 final production model: dual-terminal value + result2.xlsx.

This is the production cut of the V3.1 terminal ablation and it implements
exactly the pipeline of ``reports/Q2_最终建模思路.html``:

    Ridge load/PV forecast -> full-day residual scenarios -> exponentially
    weighted two-stage SAA -> 17-point LP-dual terminal inventory value ->
    H1 causal execution -> two-sided inventory-value accounting.

Only the adopted policy ``dual_terminal_g17`` is solved.  The four comparison
policies, the 9-point grid, the paired block bootstrap, the grid-convergence
tables and the ablation figures belong to the V3.1 study and are deliberately
out of scope here.  The one capability this module adds is the export of the
official submission workbook ``result2.xlsx``.

Units: slot flows are AC-bus kWh, SOC is kWh, price is yuan/kWh.
"""
from __future__ import annotations

from copy import copy
from dataclasses import asdict, dataclass
from pathlib import Path
import hashlib, json, os, platform, shutil, sys, time

import numpy as np
import pandas as pd
from scipy.optimize import linprog
from scipy.sparse import coo_matrix, vstack

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT / 'code'))
sys.path.insert(0, str(ROOT / 'code/q2_v2'))

from q2_planning import Study, weights_for
from dispatch import execute_day, check_dispatch
from utils import BATTERY, DT, T, INTERVALS

OUT = ROOT / 'results/q2_v32_final'
REPORT = ROOT / 'reports/Q2_V32_FINAL_RESULTS.md'
TEMPLATE = ROOT / 'C题/附件/附件5/result2.xlsx'
RESULT2 = ROOT / 'result2.xlsx'

SEED = 20260912
LO, HI, INITIAL = 1200.0, 10800.0, 6000.0
CAP = 5000.0 * DT
EPS = 1e-8
VALUE_HISTORY_DAYS = 84       # 12 complete weeks
VALUE_HISTORY_HALF = 28       # four-week recency half-life
VALUE_TOL = 1e-6              # yuan; numerical convergence only
VALUE_MAX_ITER = 50           # numerical safeguard, not a fitted parameter
GRID_SIZE = 17                # 600 kWh spacing over [1200, 10800]
POLICY = 'dual_terminal_g17'

PERIODS = {'full334': ('2025-02-01', '2025-12-31'),
           'development153': ('2025-06-01', '2025-10-31'),
           'frozen61': ('2025-11-01', '2025-12-31')}

BLOCK_LABELS = ['0:00-4:00', '4:00-8:00', '8:00-12:00',
                '12:00-16:00', '16:00-20:00', '20:00-24:00']

EMERGENCY_FLOOR = 1e-8        # slot emergency below this is written as zero
EMERGENCY_ACTIVE = 1e-7       # slot counts as an event above this
WARMUP_START, DELIVERY_START, END_DAY = 0, 31, 365


def atom(path, obj):
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + '.tmp')
    tmp.write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    os.replace(tmp, path)


def sha(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def value_at(cuts, energy):
    e = np.asarray(energy, float)
    if not cuts:
        return np.zeros_like(e, dtype=float)
    result = np.max(np.stack([a * e + b for a, b in cuts]), axis=0)
    return float(result) if result.ndim == 0 else result


def dual_supporting_cuts(grid, objectives, duals, reference=INITIAL):
    """Build exact supporting hyperplanes Q(e_j)+pi_j(e-e_j).

    For a minimization LP parameterized by the RHS initial energy, the equality
    marginal pi_j is a subgradient of the convex optimal-value function.  Hence
    each affine function is a global lower support on the physical SOC domain.
    """
    grid = np.asarray(grid, float); objectives = np.asarray(objectives, float)
    duals = np.asarray(duals, float)
    raw = [(float(pi), float(q - pi * e)) for e, q, pi in zip(grid, objectives, duals)]
    support = np.column_stack([a * grid + b for a, b in raw]).max(axis=1)
    violation = float(np.max(support - objectives))
    own_error = float(np.max(np.abs(np.array([a*e+b for e,(a,b) in zip(grid, raw)]) - objectives)))
    if violation > 1e-5 or own_error > 1e-5:
        raise AssertionError({'dual_support_violation': violation, 'own_cut_error': own_error})
    shift = float(value_at(raw, reference))
    cuts = [(a, b-shift) for a, b in raw]
    return cuts, violation, own_error


@dataclass(frozen=True)
class ScenarioConfig:
    window: int
    half_life: int
    reserve_ratio: float = .5


EARLY = ScenarioConfig(21, 14, .5)
DEPLOYED = ScenarioConfig(14, 7, .5)


class Experiment:
    def __init__(self):
        self.study = Study(ROOT)
        self.price = self.study.price
        self.dates = self.study.dates
        self.net = self.study.net
        self.fc = self.study.fc['Ridge'].copy()
        # Restore the audited cold-start forecast; never substitute actual day values.
        self.fc[:, :90] = self.study.legacy[:, :90]
        self.netfc = self.fc[0] - self.fc[1]
        self.solve_calls = 0
        self.solve_seconds = 0.0

    def config(self, day):
        # V3 rolling validation accepted 14-day/7-day on June 1 and froze it later.
        return EARLY if self.dates[day] < pd.Timestamp('2025-06-01') else DEPLOYED

    def scenarios(self, day):
        cfg = self.config(day)
        ids = np.arange(max(7, day-cfg.window), day, dtype=int)
        weights = weights_for(len(ids), cfg.half_life)
        residual = self.net[ids] - self.netfc[ids]
        scenarios = self.netfc[day][None, :] + residual  # kW
        if ids.max() >= day or not np.isfinite(scenarios).all():
            raise AssertionError('non-causal or non-finite scenario')
        return scenarios, weights, ids, cfg

    def warmup_scenarios(self, day):
        """Build causal January scenarios until the deployed rule is available."""
        if day >= 8:
            return self.scenarios(day)
        cfg = self.config(day)
        ids = np.arange(0, day, dtype=int)
        target = self.netfc[day]
        if not np.isfinite(target).all():
            target = self.study.ref[0] - self.study.ref[1]
        if len(ids) == 0:
            scenarios = target[None, :]
            weights = np.ones(1)
        else:
            weights = weights_for(len(ids), cfg.half_life)
            historical_forecast = self.netfc[ids].copy()
            missing = ~np.isfinite(historical_forecast).all(axis=1)
            historical_forecast[missing] = self.study.ref[0] - self.study.ref[1]
            residual = self.net[ids] - historical_forecast
            scenarios = target[None, :] + residual
        if (len(ids) and ids.max() >= day) or not np.isfinite(scenarios).all():
            raise AssertionError('non-causal or non-finite warmup scenario')
        return scenarios, weights, ids, cfg

    def weekday_distribution(self, origin, weekday):
        ids = np.arange(max(7, origin-VALUE_HISTORY_DAYS), origin-1, dtype=int)
        ids = ids[self.dates[ids].dayofweek == weekday]
        if len(ids) < 3:
            raise AssertionError('insufficient causal weekday history')
        weights = weights_for(len(ids), VALUE_HISTORY_HALF)
        return self.net[ids], weights, ids

    def solve(self, scenarios_kw, weights, initial_soc, cuts=None, terminal=None):
        """Solve the risk-neutral finite-scenario LP in kWh units."""
        scenarios_kw = np.atleast_2d(np.asarray(scenarios_kw, float))
        S, H = scenarios_kw.shape
        if H % T or not np.isfinite(scenarios_kw).all():
            raise ValueError('invalid scenarios')
        weights = np.asarray(weights, float); weights /= weights.sum()
        prices = np.tile(self.price, H // T)
        use_value = cuts is not None
        n0 = 4*H + S*H; n = n0 + int(use_value); theta = n0 if use_value else None

        t = np.arange(H)
        eq_row = np.r_[t, t, t, t[1:]]
        eq_col = np.r_[H+t, 2*H+t, 3*H+t, 3*H+t[:-1]]
        eq_val = np.r_[np.full(H, -BATTERY.eta_c), np.full(H, 1/BATTERY.eta_d),
                       np.ones(H), -np.ones(H-1)]
        Aeq = coo_matrix((eq_val, (eq_row, eq_col)), shape=(H, n)).tocsr()
        beq = np.zeros(H); beq[0] = float(initial_soc)

        k = np.arange(S*H); slots = np.tile(t, S)
        # q + discharge - charge + emergency >= net_power * DT
        Aub = coo_matrix((np.tile([-1., 1., -1., -1.], (S*H, 1)).ravel(),
            (np.repeat(k, 4), np.column_stack([slots, H+slots, 2*H+slots, 4*H+k]).ravel())),
            shape=(S*H, n)).tocsr()
        bub = -(scenarios_kw * DT).ravel()
        if use_value:
            rows=[]; cols=[]; vals=[]; rhs=[]
            for j, (a, b) in enumerate(cuts):
                rows.extend([j, j]); cols.extend([4*H-1, theta]); vals.extend([a, -1.]); rhs.append(-b)
            C = coo_matrix((vals, (rows, cols)), shape=(len(cuts), n)).tocsr()
            Aub = vstack([Aub, C], format='csr'); bub = np.r_[bub, rhs]

        objective = np.r_[prices, np.full(2*H, EPS), np.zeros(H),
                          (5*weights[:, None]*prices).ravel(), [1.] if use_value else []]
        energy_bounds = [(LO, HI)] * H
        if terminal is not None:
            energy_bounds[-1] = (float(terminal), float(terminal))
        bounds = ([(0, None)]*H + [(0, CAP)]*(2*H) + energy_bounds
                  + [(0, None)]*(S*H) + ([(None, None)] if use_value else []))
        started = time.perf_counter()
        ans = linprog(objective, A_ub=Aub, b_ub=bub, A_eq=Aeq, b_eq=beq,
                      bounds=bounds, method='highs', options={'presolve': True,
                      'dual_feasibility_tolerance': 1e-8, 'primal_feasibility_tolerance': 1e-8})
        if not ans.success:
            ans = linprog(objective, A_ub=Aub, b_ub=bub, A_eq=Aeq, b_eq=beq,
                          bounds=bounds, method='highs-ds', options={'presolve': False})
        elapsed = time.perf_counter() - started
        self.solve_calls += 1; self.solve_seconds += elapsed
        if not ans.success:
            raise RuntimeError(f'LP failed {ans.status}: {ans.message}')

        x = ans.x; q = x[:H].copy(); charge = x[H:2*H].copy(); discharge = x[2*H:3*H].copy()
        soc = np.r_[initial_soc, x[3*H:4*H]]
        simultaneous = np.minimum(charge, discharge).clip(0)
        if simultaneous.max(initial=0) > 1e-7:
            remove = np.minimum(charge, discharge/BATTERY.eta_c**2).clip(0)
            charge -= remove; discharge -= BATTERY.eta_c**2*remove
        emergency = np.maximum(scenarios_kw*DT + charge-discharge-q, 0)
        physical = float(prices@q + weights@(emergency@(5*prices)))
        terminal_value = float(value_at(cuts, soc[-1])) if use_value else 0.0
        eq_residual = float(np.abs(Aeq@x-beq).max())
        ineq_violation = float(max(0, np.max(Aub@x-bub)))
        if eq_residual >= 1e-5 or ineq_violation >= 1e-5:
            raise AssertionError('LP residual failure')
        return dict(q=q, charge=charge, discharge=discharge, soc=soc,
            expected_cost=physical, terminal_value=terminal_value,
            objective=physical+terminal_value, solver_objective=float(ans.fun),
            initial_dual=float(ans.eqlin.marginals[0]),
            solve_seconds=elapsed, eq_residual=eq_residual,
            ineq_violation=ineq_violation, simultaneous_kwh=float(simultaneous.sum()))


def learn_value_bank(exp, origin, grid_size=GRID_SIZE):
    grid = np.linspace(LO, HI, grid_size)
    banks = {u: [(0.0, 0.0)] for u in range(7)}
    old_values = {u: np.zeros(grid_size) for u in range(7)}
    history=[]; dual_rows=[]; check_rows=[]
    for iteration in range(1, VALUE_MAX_ITER+1):
        updated={}; new_values={}
        for weekday in range(6, -1, -1):
            next_weekday = (weekday+1) % 7
            next_cuts = updated.get(next_weekday, banks[next_weekday])
            scenarios, weights, ids = exp.weekday_distribution(origin, weekday)
            objectives=[]; duals=[]
            for energy in grid:
                solution = exp.solve(scenarios, weights, float(energy), cuts=next_cuts)
                # The dual marginal and function value must refer to exactly the
                # same perturbed solver objective; physical cost remains the
                # reporting metric outside value-function construction.
                objectives.append(solution['solver_objective']); duals.append(solution['initial_dual'])
            cuts, violation, own_error = dual_supporting_cuts(grid, objectives, duals)
            updated[weekday] = cuts; new_values[weekday] = value_at(cuts, grid)
            check_rows.append(dict(origin=str(exp.dates[origin].date()), grid_size=grid_size,
                iteration=iteration, weekday=weekday, support_violation=violation,
                own_cut_error=own_error, dual_max=max(duals), dual_min=min(duals),
                history_min_day=int(ids.min()), history_max_day=int(ids.max())))
            dual_rows.extend(dict(origin=str(exp.dates[origin].date()), grid_size=grid_size,
                iteration=iteration, weekday=weekday, initial_soc=float(e), objective=float(q),
                initial_soc_dual=float(pi), cut_slope=float(pi), cut_intercept=float(q-pi*e))
                for e,q,pi in zip(grid,objectives,duals))
        gap = max(float(np.max(np.abs(new_values[u]-old_values[u]))) for u in range(7))
        history.append(dict(origin=str(exp.dates[origin].date()), grid_size=grid_size,
                            iteration=iteration, max_value_change=gap, converged=gap<=VALUE_TOL))
        banks = updated; old_values = new_values
        if gap <= VALUE_TOL:
            break
    if history[-1]['converged'] is not True:
        raise AssertionError('relative value iteration did not converge')
    return banks, pd.DataFrame(history), pd.DataFrame(dual_rows), pd.DataFrame(check_rows)


def build_value_banks(exp):
    all_banks={}; histories=[]; duals=[]; checks=[]; curves=[]
    dense = np.linspace(LO, HI, 65)
    for month in range(2, 13):
        origin = int(np.flatnonzero(exp.dates == pd.Timestamp(2025, month, 1))[0])
        bank,h,d,c = learn_value_bank(exp, origin, GRID_SIZE)
        all_banks[month] = bank
        histories.append(h); duals.append(d); checks.append(c)
        for weekday,cuts in bank.items():
            curves.extend(dict(month=month, weekday=weekday,
                               soc=float(e), relative_future_cost=float(value_at(cuts,e))) for e in dense)
        print('value-bank month', month, 'iterations', len(h), 'gap', h.iloc[-1].max_value_change, flush=True)
    atom(OUT/'value_banks.json', {str(m): {str(u): c for u, c in b.items()} for m, b in all_banks.items()})
    pd.concat(histories, ignore_index=True).to_csv(OUT/'value_iteration.csv', index=False)
    pd.concat(duals, ignore_index=True).to_csv(OUT/'terminal_duals.csv', index=False)
    pd.concat(checks, ignore_index=True).to_csv(OUT/'dual_cut_checks.csv', index=False)
    pd.DataFrame(curves).to_csv(OUT/'terminal_value_curves.csv', index=False)
    return all_banks


def bank_for_date(exp, banks, date):
    date = pd.Timestamp(date)
    month = 12 if date.year > 2025 else int(date.month)
    return banks[month][int(date.dayofweek)]


def actual_row(exp, day, state, solution, ids, cfg):
    executed = execute_day(solution['q'], exp.net[day]*DT, state, solution['soc'], cfg.reserve_ratio)
    audit = check_dispatch(exp.net[day]*DT, solution['q'], executed['charge'], executed['discharge'],
                           executed['spill'], executed['soc'], executed['emergency'])
    if not audit['pass']:
        raise AssertionError(audit)
    planned = float(exp.price@solution['q']); emergency = float(5*exp.price@executed['emergency'])
    z = executed['emergency']; active = z > EMERGENCY_ACTIVE
    row = dict(policy=POLICY, day=day, date=str(exp.dates[day].date()),
        initial_soc=float(state), terminal_soc=float(executed['soc'][-1]),
        planned_cost=planned, emergency_cost=emergency, total_cost=planned+emergency,
        planned_kwh=float(solution['q'].sum()), emergency_kwh=float(z.sum()),
        spill_kwh=float(executed['spill'].sum()), charge_kwh=float(executed['charge'].sum()),
        discharge_kwh=float(executed['discharge'].sum()), emergency_slots=int(active.sum()),
        emergency_day=int(active.any()), events=int(np.sum(active & ~np.r_[False,active[:-1]])),
        reference_terminal_soc=float(solution['soc'][-1]), terminal_value=float(solution['terminal_value']),
        lp_objective=float(solution['objective']), solve_seconds=float(solution['solve_seconds']),
        lp_eq_residual=solution['eq_residual'], lp_ineq_violation=solution['ineq_violation'],
        balance_residual=audit['balance_max_abs'], soc_residual=audit['soc_equation_max_abs'],
        history_min_day=int(ids.min()) if len(ids) else -1,
        history_max_day=int(ids.max()) if len(ids) else -1, **asdict(cfg))
    slots = pd.DataFrame(dict(policy=POLICY, date=row['date'], slot=np.arange(T), interval=INTERVALS,
        price=exp.price, net_actual=exp.net[day]*DT, purchase=solution['q'],
        charge=executed['charge'], discharge=executed['discharge'], emergency=z,
        spill=executed['spill'], soc_start=executed['soc'][:-1], soc_end=executed['soc'][1:],
        reference_soc=solution['soc'][1:]))
    return row, slots, float(executed['soc'][-1])


def january_warmup(exp):
    """Run Jan 1--31 from the stated 6000 kWh and return the Feb 1 SOC."""
    rows=[]; state=INITIAL
    # Q3 uses the same causal cold-start boundary: value stored energy at the
    # cheapest replacement energy price, without using future January data.
    salvage = float(exp.price.min() / BATTERY.eta_d)
    cold_cuts = [(-salvage, salvage * INITIAL)]
    for day in range(WARMUP_START, DELIVERY_START):
        scenarios, weights, ids, cfg = exp.warmup_scenarios(day)
        solution = exp.solve(scenarios, weights, state, cuts=cold_cuts)
        row, _, state = actual_row(exp, day, state, solution, ids, cfg)
        row['warmup_scenario_days'] = int(len(ids))
        row['warmup_mode'] = ('attachment1_typical_day' if len(ids) == 0 else
                              'causal_residual_cold_start' if day < 8 else 'deployed_saa')
        rows.append(row)
    warmup = pd.DataFrame(rows)
    warmup.to_csv(OUT/'january_warmup_daily.csv', index=False)
    return warmup, float(state), salvage


def run_final(exp, banks):
    """Warm up January, then record the Feb 1--Dec 31 delivery period."""
    warmup, state, salvage = january_warmup(exp)
    rows=[]; slots=[]
    for day in range(DELIVERY_START, END_DAY):
        scenarios, weights, ids, cfg = exp.scenarios(day)
        cuts = bank_for_date(exp, banks, exp.dates[day]+pd.Timedelta(days=1))
        solution = exp.solve(scenarios, weights, state, cuts=cuts)
        row,slot,state = actual_row(exp, day, state, solution, ids, cfg)
        rows.append(row); slots.append(slot)
    daily = pd.DataFrame(rows); slots = pd.concat(slots, ignore_index=True)
    daily.to_csv(OUT/'daily.csv', index=False)
    slots.to_csv(OUT/'slots.csv.gz', index=False, compression='gzip')
    print('policy', POLICY, 'completed', len(daily), 'days', flush=True)
    return warmup, daily, slots, salvage


def inventory_boundary(exp, banks, first_date, first_soc, last_date, last_soc):
    start_value = float(value_at(bank_for_date(exp, banks, first_date), first_soc))
    end_value = float(value_at(bank_for_date(exp, banks, pd.Timestamp(last_date)+pd.Timedelta(days=1)), last_soc))
    return start_value, end_value, end_value-start_value


def summarize(exp, daily, banks):
    data=daily.copy(); data['date']=pd.to_datetime(data.date)
    summaries=[]
    for period,(a,b) in PERIODS.items():
        q=data[data.date.between(a,b)].sort_values('date')
        start_value,end_value,adjustment = inventory_boundary(exp,banks,a,q.iloc[0].initial_soc,
                                                              b,q.iloc[-1].terminal_soc)
        summaries.append(dict(policy=POLICY,period=period,days=len(q),
            planned_cost=q.planned_cost.sum(),emergency_cost=q.emergency_cost.sum(),
            raw_total_cost=q.total_cost.sum(),start_inventory_value=start_value,
            end_inventory_value=end_value,inventory_adjustment=adjustment,
            adjusted_total_cost=q.total_cost.sum()+adjustment,planned_kwh=q.planned_kwh.sum(),
            emergency_kwh=q.emergency_kwh.sum(),spill_kwh=q.spill_kwh.sum(),
            emergency_slots=q.emergency_slots.sum(),emergency_day_frequency=q.emergency_day.mean(),
            max_daily_cost=q.total_cost.max(),start_soc=q.iloc[0].initial_soc,
            end_soc=q.iloc[-1].terminal_soc))
    summary=pd.DataFrame(summaries); summary.to_csv(OUT/'summary.csv',index=False)
    data['month']=data.date.dt.month
    data.groupby(['month'])[['planned_cost','emergency_cost','total_cost','emergency_kwh','spill_kwh']].sum().reset_index().to_csv(OUT/'monthly.csv',index=False)
    return summary


def verify(exp,warmup,daily):
    dual_checks=pd.read_csv(OUT/'dual_cut_checks.csv');vi=pd.read_csv(OUT/'value_iteration.csv')
    tests=[]
    def add(name,value,tolerance):
        passed=bool(value<=tolerance);tests.append(dict(test=name,value=float(value),tolerance=float(tolerance),passed=passed))
        if not passed:raise AssertionError((name,value,tolerance))
    add('causality',max(0,float((daily.history_max_day-daily.day+1).max())),0)
    observed_warmup = warmup[warmup.warmup_scenario_days > 0]
    add('january_warmup_causality',
        max(0, float((observed_warmup.history_max_day-observed_warmup.day+1).max())), 0)
    add('january_initial_soc', abs(float(warmup.iloc[0].initial_soc)-INITIAL), 1e-9)
    add('february_inherits_january_soc',
        abs(float(daily.iloc[0].initial_soc)-float(warmup.iloc[-1].terminal_soc)), 1e-9)
    add('balance',float(daily.balance_residual.max()),1e-7)
    add('soc_recursion',float(daily.soc_residual.max()),1e-7)
    add('lp_eq',float(daily.lp_eq_residual.max()),1e-5)
    add('lp_ineq',float(daily.lp_ineq_violation.max()),1e-5)
    add('dual_support',max(0,float(dual_checks.support_violation.max())),1e-5)
    add('dual_own_cut',float(dual_checks.own_cut_error.max()),1e-5)
    add('value_iteration_convergence',float((~vi.groupby('origin').tail(1).converged).sum()),0)
    add('january_soc_continuity',
        float(np.max(np.abs(warmup.initial_soc.iloc[1:].to_numpy()
                            -warmup.terminal_soc.iloc[:-1].to_numpy()))),1e-7)
    add('soc_continuity_'+POLICY,
        float(np.max(np.abs(daily.initial_soc.iloc[1:].to_numpy()-daily.terminal_soc.iloc[:-1].to_numpy()))),1e-7)
    result={'all_pass':all(x['passed'] for x in tests),'tests':tests,
            'solve_calls':exp.solve_calls,'solve_seconds':exp.solve_seconds}
    atom(OUT/'verification.json',result);return result


def clock_label(total_minutes):
    """Slot boundary label, e.g. 210 -> '3:30', 1440 -> '24:00'."""
    return '24:00' if total_minutes == 1440 else f'{total_minutes//60}:{total_minutes%60:02d}'


def merge_events(z):
    """Merge contiguous active slots into one interval, as in 题面表 4."""
    active = np.asarray(z, float) > EMERGENCY_ACTIVE
    events=[]; i=0
    while i < len(active):
        if not active[i]:
            i += 1; continue
        j = i
        while j+1 < len(active) and active[j+1]:
            j += 1
        events.append(dict(time_period=f'{clock_label(10*i)}-{clock_label(10*(j+1))}',
                           emergency_kwh=float(np.asarray(z,float)[i:j+1].sum()),
                           slot_from=i, slot_to=j))
        i = j+1
    return events


def slot_matrix(slots, order, column):
    return slots.pivot(index='date', columns='slot', values=column).loc[order].to_numpy(float)


def build_result2_frame(slots):
    """Shape the executed schedule into the three workbook tables."""
    order = sorted(slots.date.unique())
    q = slot_matrix(slots, order, 'purchase')
    charge = slot_matrix(slots, order, 'charge')
    discharge = slot_matrix(slots, order, 'discharge')
    soc_start = slot_matrix(slots, order, 'soc_start')
    soc_end = slot_matrix(slots, order, 'soc_end')
    price = slot_matrix(slots, order, 'price')[0]
    z = slot_matrix(slots, order, 'emergency')
    events = []
    for pos, date in enumerate(order):
        row = z[pos].copy(); row[row < EMERGENCY_FLOOR] = 0.0
        for e in merge_events(row):
            events.append(dict(date=str(date), time_period=e['time_period'],
                               emergency_kwh=e['emergency_kwh'],
                               slot_from=e['slot_from'], slot_to=e['slot_to']))
    frame = pd.DataFrame(events, columns=['date', 'time_period', 'emergency_kwh',
                                          'slot_from', 'slot_to'])
    return dict(order=order, q=q, charge=charge, discharge=discharge, soc_start=soc_start,
                soc_end=soc_end, price=price, emergency=z, events=frame)


def write_result2(slots, out=RESULT2, template=TEMPLATE):
    """Fill the official submission workbook from the executed schedule."""
    from openpyxl import load_workbook
    if not template.exists():
        raise FileNotFoundError(f'missing template: {template}')
    d = build_result2_frame(slots)
    order, q, price = d['order'], d['q'], d['price']
    charge, discharge = d['charge'], d['discharge']
    soc_start, soc_end, events = d['soc_start'], d['soc_end'], d['events']
    days = len(order)

    shutil.copyfile(template, out)
    wb = load_workbook(out)

    ws_plan = wb['计划购电量']
    if ws_plan.max_row < 1+days or ws_plan.max_column < 147:
        raise AssertionError('计划购电量 template too small')
    for pos, date in enumerate(order):
        row = pos + 2
        ws_plan.cell(row, 1, pd.Timestamp(date).to_pydatetime())
        ws_plan.cell(row, 1).number_format = 'yyyy/m/d'
        for t in range(T):
            ws_plan.cell(row, 2+t, float(q[pos, t]))
            ws_plan.cell(row, 2+t).number_format = '0.0000'
        ws_plan.cell(row, 146, float(q[pos].sum()))
        ws_plan.cell(row, 147, float(price @ q[pos]))
        ws_plan.cell(row, 146).number_format = ws_plan.cell(row, 147).number_format = '0.0000'

    ws_st = wb['充放电量']
    block_style = [[copy(ws_st.cell(r, c)._style) for c in range(1, 7)] for r in range(2, 8)]
    block_height = [ws_st.row_dimensions[r].height for r in range(2, 8)]
    ws_st.delete_rows(2, ws_st.max_row - 1)
    for pos, date in enumerate(order):
        for b in range(6):
            row = 2 + 6*pos + b
            for c in range(1, 7):
                ws_st.cell(row, c)._style = copy(block_style[b][c-1])
            ws_st.row_dimensions[row].height = block_height[b]
            lo, hi = 24*b, 24*(b+1)
            ws_st.cell(row, 1, pd.Timestamp(date).to_pydatetime() if b == 0 else None)
            ws_st.cell(row, 2, BLOCK_LABELS[b])
            ws_st.cell(row, 3, float(charge[pos, lo:hi].sum()))
            ws_st.cell(row, 4, float(discharge[pos, lo:hi].sum()))
            if b == 0:
                ws_st.cell(row, 5, '0:00'); ws_st.cell(row, 6, float(soc_start[pos, 0]))
                ws_st.cell(row, 1).number_format = 'yyyy/m/d'
            elif b == 1:
                ws_st.cell(row, 5, '24:00'); ws_st.cell(row, 6, float(soc_end[pos, -1]))
            for c in (3, 4, 6):
                ws_st.cell(row, c).number_format = '0.0000'

    ws_ev = wb['紧急购电量']
    ev_style = {kind: [copy(ws_ev.cell(src, c)._style) for c in range(1, 4)]
                for kind, src in [('first', 2), ('middle', 3), ('last', 4)]}
    ws_ev.delete_rows(2, ws_ev.max_row - 1)
    row = 2
    for date, group in events.groupby('date', sort=True):
        group = group.reset_index(drop=True)
        for i, event in group.iterrows():
            kind = 'first' if i == 0 else ('last' if i == len(group)-1 else 'middle')
            for c in range(1, 4):
                ws_ev.cell(row, c)._style = copy(ev_style[kind][c-1])
            ws_ev.cell(row, 1, pd.Timestamp(date).to_pydatetime() if i == 0 else None)
            ws_ev.cell(row, 2, event['time_period'])
            ws_ev.cell(row, 3, float(event['emergency_kwh']))
            if i == 0:
                ws_ev.cell(row, 1).number_format = 'yyyy/m/d'
            ws_ev.cell(row, 3).number_format = '0.0000'
            row += 1

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(out)
    events.to_csv(OUT/'result2_emergency_events.csv', index=False)

    audit = audit_result2(out, d)
    atom(OUT/'result2_export_audit.json', audit)
    if not audit['passed']:
        raise AssertionError(audit)
    print('result2 export audit:', json.dumps(audit['checks'], ensure_ascii=False), flush=True)
    return audit


def audit_result2(path, d):
    """Read the written workbook back and compare cell by cell."""
    from openpyxl import load_workbook
    q, charge = d['q'], d['charge']
    discharge, soc_start, soc_end = d['discharge'], d['soc_start'], d['soc_end']
    price, events, z = d['price'], d['events'], d['emergency']
    days = q.shape[0]

    wb = load_workbook(path, data_only=False)
    ws_plan, ws_st, ws_ev = wb['计划购电量'], wb['充放电量'], wb['紧急购电量']
    plan = np.array([[ws_plan.cell(r, c).value for c in range(2, 146)] for r in range(2, 2+days)], float)
    totals = np.array([[ws_plan.cell(r, c).value for c in (146, 147)] for r in range(2, 2+days)], float)
    block_c = np.array([ws_st.cell(r, 3).value for r in range(2, 2+6*days)], float)
    block_d = np.array([ws_st.cell(r, 4).value for r in range(2, 2+6*days)], float)
    soc_head = np.array([float(ws_st.cell(2+6*i, 6).value) for i in range(days)])
    soc_tail = np.array([float(ws_st.cell(3+6*i, 6).value) for i in range(days)])
    shape_st = (ws_st.max_row, ws_st.max_column)
    shape_ev = (ws_ev.max_row, ws_ev.max_column)
    ev = np.array([ws_ev.cell(r, 3).value for r in range(2, 2+len(events))], float) if len(events) else np.zeros(0)
    formula_errors = [f'{ws.title}!{c.coordinate}:{c.value}'
                      for ws in wb.worksheets for row in ws.iter_rows() for c in row
                      if isinstance(c.value, str)
                      and c.value.startswith(('#REF!', '#DIV/0!', '#VALUE!', '#N/A', '#NAME?'))]
    wb.close()

    exp_c = np.array([charge[i, 24*b:24*(b+1)].sum() for i in range(days) for b in range(6)])
    exp_d = np.array([discharge[i, 24*b:24*(b+1)].sum() for i in range(days) for b in range(6)])
    written_energy = float(ev.sum()) if len(events) else 0.0

    checks = {
        'plan_shape_ok': bool((ws_plan.max_row, ws_plan.max_column) == (335, 147)),
        'storage_shape_ok': bool(shape_st == (1+6*days, 6)),
        'event_shape_ok': bool(shape_ev == (1+len(events), 3)),
        'plan_max_abs_error_kwh': float(np.max(np.abs(plan - q))),
        'total_kwh_max_abs_error': float(np.max(np.abs(totals[:, 0] - q.sum(axis=1)))),
        'total_cost_max_abs_error_yuan': float(np.max(np.abs(totals[:, 1] - q @ price))),
        'storage_charge_max_abs_error_kwh': float(np.max(np.abs(block_c - exp_c))),
        'storage_discharge_max_abs_error_kwh': float(np.max(np.abs(block_d - exp_d))),
        'soc_start_max_abs_error_kwh': float(np.max(np.abs(soc_head - soc_start[:, 0]))),
        'soc_end_max_abs_error_kwh': float(np.max(np.abs(soc_tail - soc_end[:, -1]))),
        'event_energy_abs_error_kwh': float(abs(written_energy - float(z.sum()))),
        'formula_error_count': len(formula_errors),
    }
    limits = {'plan_max_abs_error_kwh': 1e-8, 'total_kwh_max_abs_error': 1e-8,
              'total_cost_max_abs_error_yuan': 1e-6,
              'storage_charge_max_abs_error_kwh': 1e-8, 'storage_discharge_max_abs_error_kwh': 1e-8,
              'soc_start_max_abs_error_kwh': 1e-8, 'soc_end_max_abs_error_kwh': 1e-8,
              'event_energy_abs_error_kwh': 1e-5, 'formula_error_count': 0}
    checks['pass'] = bool(checks['plan_shape_ok'] and checks['storage_shape_ok']
                          and checks['event_shape_ok']
                          and all(checks[k] <= v for k, v in limits.items()))
    return dict(path=str(path), days=int(days), events=int(len(events)),
                planned_kwh=float(q.sum()), planned_cost=float(totals[:, 1].sum()),
                written_emergency_kwh=written_energy, model_emergency_kwh=float(z.sum()),
                passed=checks['pass'], checks=checks)


def full():
    if OUT.exists(): shutil.rmtree(OUT)
    OUT.mkdir(parents=True)
    started = time.time()
    exp = Experiment()
    banks = build_value_banks(exp)
    warmup, daily, slots, warmup_salvage = run_final(exp, banks)
    summary = summarize(exp, daily, banks)
    verification = verify(exp, warmup, daily)
    export = write_result2(slots)
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text('# Q2 V3.2 最终版运行结果\n\n本报告由本次 Colab 输出自动生成。\n\n'
        + '## 费用（双边库存价值修正）\n\n' + summary.to_markdown(index=False)
        + '\n\n## 验收\n\n' + pd.DataFrame(verification['tests']).to_markdown(index=False)
        + '\n\n## result2.xlsx 回读校验\n\n'
        + pd.DataFrame([{k: v for k, v in export['checks'].items()}]).T.rename(columns={0: 'value'}).to_markdown()
        + '\n', encoding='utf-8')
    manifest = dict(complete=True, started_at=pd.Timestamp.fromtimestamp(started, tz='UTC').isoformat(),
        completed_at=pd.Timestamp.now(tz='UTC').isoformat(), elapsed_seconds=time.time()-started,
        solve_calls=exp.solve_calls, solve_seconds=exp.solve_seconds, python=platform.python_version(),
        cpu_count=os.cpu_count(), notebook='code/problem2_v32_final.ipynb',
        source_sha256=sha(__file__),
        prediction_sha256=sha(ROOT/'results/problem2_forecast_ablation_predictions.csv.gz'),
        grid_size=GRID_SIZE, policy=POLICY, result2=str(RESULT2),
        initial_jan1_kwh=INITIAL, initial_feb1_kwh=float(daily.iloc[0].initial_soc),
        january_warmup_days=len(warmup),
        january_cold_start=('Attachment 1 typical-day forecast while frozen Ridge is unavailable; '
                            'causal completed-day residuals on Jan 2-8; deployed SAA from Jan 9'),
        january_terminal_inventory_value_yuan_per_kwh=warmup_salvage,
        delivery_period='2025-02-01/2025-12-31',
        all_checks_pass=bool(verification['all_pass'] and export['passed']), fresh_run=True)
    atom(OUT/'run_manifest.json', manifest)
    print(json.dumps(manifest, ensure_ascii=False, indent=2), flush=True)


if __name__ == '__main__':
    full()
