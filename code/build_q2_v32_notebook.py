"""Build the self-contained Q2 V3.2 final Colab notebook."""
from pathlib import Path
import nbformat as nbf

ROOT=Path(__file__).resolve().parents[1]
MODULE=ROOT/'code/q2_v32/q2_final.py'
OUT=ROOT/'code/problem2_v32_final.ipynb'
nb=nbf.v4.new_notebook()
nb.metadata.update({'kernelspec':{'display_name':'Python 3','language':'python','name':'python3'},
    'language_info':{'name':'python','version':'3'},
    'colab':{'name':OUT.name,'provenance':[],'gpuType':'L4'}})

def md(x):nb.cells.append(nbf.v4.new_markdown_cell(x.strip()))
def code(x):nb.cells.append(nbf.v4.new_code_cell(x.strip()))

md(r'''
# C题问题二 最终版：对偶终端价值的因果滚动两阶段 SAA

本笔记本是论文问题二最终方案的落地版本，实现 `reports/Q2_最终建模思路.html` 的完整口径：

预测中心（Ridge 负荷/光伏分解预测）→ 完整日残差指数加权场景 → 风险中性两阶段 SAA → 17 点 LP 对偶终端库存价值 → H1 逐时因果执行 → 双边库存价值评价。

相对 V3.1 终端消融实验，本版只求解最终采用的 `dual_terminal_g17` 一条策略。其余四条对照策略、9 点网格、配对 Bootstrap、网格收敛表与消融图属于研究工作，不在本版范围内；本版新增的唯一能力是把逐槽执行结果写入正式提交工作簿 `result2.xlsx`。

`result2.xlsx` 会被本次运行覆盖，覆盖前请确认已有备份。
''')

md(r'''
## 1. 信息集、单位与时间口径

第 \(d\) 天第 \(t\) 个 10 分钟时槽的净负荷功率为

\[
N_{d,t}=L_{d,t}-P_{d,t}\quad(\mathrm{kW}),
\qquad \Delta t=\frac{1}{6}\ \mathrm{h},
\]

进入模型的时槽净电量为 \(n_{d,t}=N_{d,t}\Delta t\)（kWh）。代码中 \(q_t,c_t,v_t,z_{s,t}\) 均表示一个时槽内的电量，\(E_t\) 表示时槽末 SOC，电价 \(p_t\) 单位为元/kWh。日期 \(d\) 的决策只使用 \(i<d\) 的历史数据。

附件时间标签为**右端点**（`0:10` 对应 `00:00-00:10`，`0:00+1` 对应 `23:50-24:00`）。附件 5 模板 144 个区间名整体后移一格，本版**不改表头**，按位置对应写入：模板第 \(k\) 个数据格 ↔ 附件第 \(k\) 列 ↔ 模型第 \(k\) 槽。
''')

code(r'''
from pathlib import Path
import os,sys,json
import numpy as np,pandas as pd
from IPython.display import display,Markdown
PROJECT=Path('/content/CUMCM_2026_Last_Dance') if Path('/content/CUMCM_2026_Last_Dance').exists() else Path.cwd()
os.chdir(PROJECT)
for extra in ('code','code/q2_v2','code/q2_v32'):
    sys.path.insert(0,str(PROJECT/extra))
print('project =',PROJECT)
print('python =',sys.version.split()[0])
''')

md(r'''
## 2. 冻结的预测与残差场景

复用既有 Ridge 对负荷、光伏的预测，构造点预测净负荷 \(\widehat N_{d,t}=\widehat L_{d,t}-\widehat P_{d,t}\)。历史完整日残差为 \(\varepsilon_{i,t}=N_{i,t}-\widehat N_{i,t}\)，第 \(i\) 条场景为

\[
N_{d,t}^{(i)}=\widehat N_{d,t}+\varepsilon_{i,t}.
\]

2—5 月沿用 21 日窗口、14 日半衰期；6 月 1 日起冻结为 14 日窗口、7 日半衰期。场景权重为

\[
w_{d,i}=\frac{2^{-(d-i)/\tau}}{\sum_j 2^{-(d-j)/\tau}}.
\]

执行储备系数保持已验证的 \(\rho=0.5\)，本版不重新搜索其他规划参数。1 月作为历史预热期，2025-02-01 0:00 实际 SOC 取 6000 kWh（题设）。
''')

code(r'''
import q2_final as M
preview=M.Experiment()
audit=pd.DataFrame({
    '检查项':['日期数','每日时槽','负荷有限','光伏有限','日期唯一','日期递增','预测文件哈希'],
    '结果':[len(preview.dates),preview.net.shape[1],np.isfinite(preview.study.load).all(),
          np.isfinite(preview.study.pv).all(),preview.dates.is_unique,preview.dates.is_monotonic_increasing,
          M.sha(PROJECT/'results/problem2_forecast_ablation_predictions.csv.gz')]})
display(audit)
print('决策日区间:',preview.dates[31].date(),'→',preview.dates[364].date(),'共',364-31+1,'天')
''')

md(r'''
## 3. 风险中性有限场景线性规划

决策变量：计划购电量 \(q_t\)、参考充电量 \(c_t\)、参考放电量 \(v_t\)、SOC \(E_t\)、场景紧急购电量 \(z_{s,t}\) 与终端未来成本 \(\theta\)。

\[
\min\quad \sum_{t=1}^{144}p_tq_t+\sum_{s=1}^{S}w_s\sum_{t=1}^{144}5p_tz_{s,t}+\theta,
\]

\[
q_t+v_t-c_t+z_{s,t}\ge N_{d,t}^{(s)}\Delta t,\qquad
E_t=E_{t-1}+\eta_c c_t-\frac{v_t}{\eta_d},
\]

\[
0\le c_t,v_t\le P_{\max}\Delta t=833.333\ \mathrm{kWh},\quad
1200\le E_t\le 10800,\quad q_t,z_{s,t}\ge0 .
\]

目标与约束均为线性，因此每个有限场景问题都是 LP，求解器给出的原始—对偶证书对应有限场景问题的全局最优解。该结论不外推为未知真实分布下的全局最优。
''')

md(r'''
## 4. 由 LP 对偶变量生成终端库存价值

取消每日硬终端等式。在 SOC 网格点 \(e_j\) 上求 Bellman 子问题得到最优值 \(Q_u(e_j)\)；初始 SOC 所在等式右端的对偶变量

\[
\pi_j=\frac{\partial Q_u(e_j)}{\partial e_j}
\]

是凸最优值函数的一个次梯度，因此支持切平面为

\[
\ell_j(E)=Q_u(e_j)+\pi_j(E-e_j),
\qquad
\widehat V_u(E)=\max_j\ell_j(E)-\max_j\ell_j(6000).
\]

6000 kWh 仅是加法零点，不要求实际或参考日末 SOC 等于 6000。本版使用 **17 个等距网格点**覆盖 \(([1200,10800])\) kWh，间距 600 kWh；价值迭代停止容差 \(10^{-6}\) 元。每个月初用此前 84 天（12 个完整周）、同星期完整日、28 日半衰期加权估计各星期净负荷分布。
''')

md(r'''
## 5. H1 逐时因果执行与双边库存价值

日前模型只输出购电 \(q_t\) 与参考 SOC 路径 \(E_t^{\rm ref}\)。执行器在每个时槽只读取当前实际净负荷、当前实际 SOC、日前购电量与当前参考 SOC：购电有余额时优先充电并弃余电；购电不足时先按动态储备线

\[
R_t=1200+\rho\left(E_t^{\rm ref}-1200\right),\qquad \rho=0.5
\]

约束放电，残余缺口按 5 倍电价紧急购电。实际 SOC 按递推式跨日连续继承，不访问任何未来真实值。

对评价区间 \(([a,b])\)，双边库存价值修正为

\[
C_{[a,b]}^{\rm adj}=C_{[a,b]}^{\rm raw}+\widehat V_{b+1}(E_b)-\widehat V_a(E_{a-1}).
\]
''')

md(r'''
## 6. `result2.xlsx` 填写口径

工作簿由附件 5 模板复制后原地填充，保留原有样式与列宽：

| 工作表 | 内容 | 规模 |
| --- | --- | --- |
| 计划购电量 | 每日 0:00 制定的 144 个 10 分钟计划购电量，另附「全天购电量」「全天购电费」 | 334 行 × 147 列 |
| 充放电量 | 六个 4 小时时段的最终实际充、放电量，以及 0:00 与 24:00 储电量 | 334 × 6 行 |
| 紧急购电量 | 实际发生的紧急购电事件，连续时槽合并为同一时段（题面表 4 格式） | 事件行数 |

储能侧写的是**实际执行结果**，不是日前参考轨迹；紧急购电同样只记录实际发生量。
''')

code(r'''
# 以下是最终版的完整可执行源码。
__file__=str(PROJECT/'code/q2_v32/q2_final.py')
'''+MODULE.read_text(encoding='utf-8').replace('from __future__ import annotations\n','')
   .replace("\nif __name__ == '__main__':\n    full()\n",'\n'))

md('''## 7. 从头运行最终版

本单元删除并重建 `results/q2_v32_final/`，重建 `result2.xlsx`。预测文件只读，其他实验的结果不作为本次输入。''')
code('full()')

md('''## 8. 费用与分期结果（双边库存价值修正）''')
code(r'''
summary=pd.read_csv(OUT/'summary.csv')
display(summary)
''')

md('''## 9. 约束、因果性与数值验收''')
code(r'''
verification=json.loads((OUT/'verification.json').read_text())
display(pd.DataFrame(verification['tests']))
dual=pd.read_csv(OUT/'dual_cut_checks.csv')
display(dual[['support_violation','own_cut_error','dual_min','dual_max']].agg(['min','max']))
vi=pd.read_csv(OUT/'value_iteration.csv')
display(vi.groupby('origin').tail(1).reset_index(drop=True))
assert verification['all_pass']
''')

md('''## 10. `result2.xlsx` 回读校验

写入后重新打开工作簿逐格比对：计划购电量、日总计、六个时段的充放电量、0:00/24:00 储电量、事件条数与事件电量，并扫描公式错误。''')
code(r'''
export=json.loads((OUT/'result2_export_audit.json').read_text())
print('days =',export['days'],'| events =',export['events'],'| passed =',export['passed'])
display(pd.Series(export['checks'],name='value').to_frame())
display(pd.DataFrame([{
    '计划购电量/kWh':export['planned_kwh'],
    '全天购电费/元':export['planned_cost'],
    '紧急购电量/kWh':export['written_emergency_kwh']}]))
from openpyxl import load_workbook
wb=load_workbook(PROJECT/'result2.xlsx')
display(pd.DataFrame([{'工作表':ws.title,'行数':ws.max_row,'列数':ws.max_column} for ws in wb.worksheets]))
ws=wb['计划购电量']
display(pd.DataFrame([[('' if c.value is None else c.value) for c in ws[r]][:4] for r in (2,3,335)],
                     columns=['日期','第1槽','第2槽','第3槽']))
wb.close()
''')

md('''## 11. 运行清单与结论''')
code(r'''
manifest=json.loads((OUT/'run_manifest.json').read_text())
print(json.dumps(manifest,ensure_ascii=False,indent=2))
full_rows=summary.set_index('period')
display(Markdown(f"""最终策略为 **{manifest['policy']}**（{manifest['grid_size']} 点对偶终端价值网格）。全年 334 天原始实付总费用
**{full_rows.loc['full334','raw_total_cost']:,.2f} 元**，双边库存价值修正后 **{full_rows.loc['full334','adjusted_total_cost']:,.2f} 元**；
计划购电费 {full_rows.loc['full334','planned_cost']:,.2f} 元，紧急购电费 {full_rows.loc['full334','emergency_cost']:,.2f} 元，
紧急购电量 {full_rows.loc['full334','emergency_kwh']:,.2f} kWh，期末实际 SOC {full_rows.loc['full334','end_soc']:,.2f} kWh。
本次共求解 {manifest['solve_calls']:,} 个 LP，纯求解 {manifest['solve_seconds']:.2f} s。
`result2.xlsx` 已按位置对应写入并通过逐格回读校验（{export['events']} 条紧急购电事件）。"""))
''')

nbf.write(nb,OUT)
print(OUT)
